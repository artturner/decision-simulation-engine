#!/usr/bin/env python
"""Per-student Unit 2 progress + access-code emails, as a mail-merge file.

Builds one row per student: their ETAMU email (from a D2L export), their
personal access code, a short Unit 2 summary from the same live data and
matching rules as checklist_unit2.py, and a ready-to-send subject + HTML
body. Nothing is sent from here — send from your own ETAMU mailbox (Power
Automate "Send an email (V2)" per row, or Word mail merge with desktop
Outlook) so the mail comes from a trusted university sender.

    python fetch_unit2.py                         # refresh live data first
    python email_merge_unit2.py --emails "C:\\Users\\arttu\\Downloads\\2026_fall_emails.csv"
    python email_merge_unit2.py --emails ... --only-unclaimed
    python email_merge_unit2.py --emails ... --test-to you@etamu.edu   # 3-row test file

Outputs (OneDrive 2026 Fall/email-merge/):
  unit2_email_merge_<date>.xlsx  (Excel table "Merge" — for Power Automate)
  unit2_email_merge_<date>.csv   (same rows — for Word mail merge)
  unit2_email_preview_<date>.html (first few emails rendered, for review)

The files contain access codes and grades: keep them in OneDrive, delete
them after sending.
"""
from __future__ import annotations

import argparse
import csv
import html
import os
import sys
from datetime import date
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import checklist_unit2 as ck  # noqa: E402  (loads unit2_raw + rosetta; renders nothing)

OUT_DIR = Path(os.environ.get("EMAIL_MERGE_OUT_DIR")
               or r"C:\Users\arttu\OneDrive - Grand Prairie ISD\2026 Fall\email-merge")
DASHBOARD = "https://scenarios.cruxlabs.academy/me"
DASHBOARD_SHORT = "scenarios.cruxlabs.academy/me"
FRQ_DUE = date(2026, 10, 8)
# Every other Unit 2 item's pacing-guide target is on or before Oct 2.
OTHER_DUE = date(2026, 10, 2)
SIGNATURE = "Arthur Turner<br>PSCI 2305 American Government"
SUBJECT = "PSCI 2305: your Unit 2 progress and personal access code"


def esc(s) -> str:
    return html.escape(str(s if s is not None else ""))


def progress(rec: dict, today: date) -> dict:
    """The 19 Unit 2 work items, counted exactly as checklist_unit2.sheet does."""
    items: list[tuple[str, bool, bool]] = []  # (label, done, is_frq)
    for title in ck.VIDEOS:
        items.append((f"Watch {title}", rec["videos"].get(title) == "completed", False))
    for key, title in ck.SCENARIOS.items():
        s = rec["scenarios"].get(key, {"status": "not_started", "reflected": False})
        items.append((f"Play {title}", s["status"] == "completed", False))
        items.append((f"Submit the {title} reflection", bool(s["reflected"]), False))
    es = rec["essay"] or "not_started"
    items.append(("FRQ 2", es in ("graded", "accepted"), True))
    done = sum(1 for _, d, _ in items if d)
    past_due = [label for label, d, frq in items
                if not d and not frq and today > OTHER_DUE]
    frq_status = {"graded": "submitted", "accepted": "submitted",
                  "draft": "started (draft not yet submitted)"}.get(es, "not started")
    return {"done": done, "total": len(items), "past_due": past_due,
            "frq_done": es in ("graded", "accepted"), "frq_status": frq_status}


def body_html(first: str, code: str | None, claimed: bool, p: dict, today: date) -> str:
    days = (FRQ_DUE - today).days
    when = "today" if days == 0 else "tomorrow" if days == 1 else f"{FRQ_DUE:%A}, {FRQ_DUE:%b} {FRQ_DUE.day}"
    lines = [f"<p>Hi {esc(first)},</p>",
             f"<p>Here's where you stand on Unit 2 as of {today:%A}, {today:%b} {today.day}:</p><ul>",
             f"<li><b>{p['done']} of {p['total']}</b> Unit 2 items done"
             + (f", <b>{len(p['past_due'])} past due</b>" if p["past_due"] else "") + "</li>"]
    if p["frq_done"]:
        lines.append("<li>FRQ 2: submitted. Nice work.</li>")
    else:
        lines.append(f"<li>FRQ 2 is due <b>{when}</b>. Yours is {esc(p['frq_status'])}.</li>")
    lines.append("</ul>")
    if p["past_due"]:
        shown = p["past_due"][:3]
        more = len(p["past_due"]) - len(shown)
        lines.append("<p>Best place to start: " + "; ".join(esc(x) for x in shown)
                     + (f"; and {more} more" if more else "")
                     + ". Late work still counts (-1% per day through Dec 11).</p>")
    if claimed:
        lines.append(f"<p>You're already set up with your access code. Your live dashboard shows every "
                     f"grade, what's due, and your best next move: "
                     f'<a href="{DASHBOARD}">{DASHBOARD_SHORT}</a></p>')
    elif code:
        lines.append(f"<p>Your personal access code is <b style=\"font-family:Consolas,monospace;"
                     f"font-size:16px;letter-spacing:2px\">{esc(code)}</b></p>"
                     f"<p>Enter it at <a href=\"{DASHBOARD}\">{DASHBOARD_SHORT}</a> to unlock your live "
                     "dashboard: every grade, what's due, and your best next move. Soon this code will "
                     "also be required to open your scenario and essay work, so set it up now. "
                     "Keep it private; it's tied to your grades.</p>")
    lines.append("<p>Questions? Just reply to this email.</p>")
    lines.append(f"<p>{SIGNATURE}</p>")
    return "\n".join(lines)


CODE_ONLY_SUBJECT = "PSCI 2305: your personal access code"


def code_only_html(first: str, code: str | None, claimed: bool) -> str:
    """No grade information at all: just the code and where to use it."""
    lines = [f"<p>Hi {esc(first)},</p>"]
    if claimed:
        lines.append("<p>You're already set up with your PSCI 2305 access code. Your live dashboard "
                     "shows every grade, what's due, and your best next move: "
                     f'<a href="{DASHBOARD}">{DASHBOARD_SHORT}</a></p>')
    else:
        lines.append("<p>As I mentioned in class, here is your personal access code for our course "
                     "websites:</p>"
                     f"<p><b style=\"font-family:Consolas,monospace;font-size:18px;letter-spacing:2px\">"
                     f"{esc(code)}</b></p>"
                     f"<p>Enter it at <a href=\"{DASHBOARD}\">{DASHBOARD_SHORT}</a> to unlock your live "
                     "dashboard: every grade, what's due, and your best next move. Soon this code will "
                     "also be required to open your scenario and essay work, so set it up now.</p>"
                     "<p>Keep it private; it's tied to your grades. If you didn't expect this email, "
                     "ask me in class before using the code.</p>")
    lines.append("<p>Questions? Reply to this email or ask me in class.</p>")
    lines.append(f"<p>{SIGNATURE}</p>")
    return "\n".join(lines)


def to_text(h: str) -> str:
    import re
    t = re.sub(r"<li>", "- ", h)
    t = re.sub(r"<br>|</p>|</li>|</ul>", "\n", t)
    t = re.sub(r'<a href="([^"]+)">[^<]*</a>', r"\1", t)
    t = re.sub(r"<[^>]+>", "", t)
    return html.unescape(re.sub(r"\n{3,}", "\n\n", t)).strip()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--emails", type=Path, required=True, help="D2L export with OrgDefinedId + Email")
    ap.add_argument("--only-unclaimed", action="store_true", help="only students who haven't claimed a code")
    ap.add_argument("--test-to", help="write a 3-row TEST file with every Email replaced by this address")
    ap.add_argument("--as-of", type=date.fromisoformat, default=date.today())
    ap.add_argument("--code-only", action="store_true",
                    help="email only the access code + dashboard link, no grade information")
    args = ap.parse_args()

    with open(args.emails, encoding="utf-8-sig", newline="") as f:
        emails = {r["OrgDefinedId"].lstrip("#").strip(): r["Email"].strip()
                  for r in csv.DictReader(f) if (r.get("Email") or "").strip()}

    by_oid = {rec["oid"]: rec for rec in ck.students.values() if rec.get("oid")}
    rows, no_data, no_code = [], [], []
    for oid, email in emails.items():
        rec = by_oid.get(oid)
        if rec is None:
            no_data.append(oid)
            continue
        claim = rec.get("claim")
        if not claim:
            no_code.append(f"{rec['first']} {rec['last']}")
        claimed = bool(claim and claim["claimed"])
        if args.only_unclaimed and claimed:
            continue
        p = progress(rec, args.as_of)
        if args.code_only:
            if not claim:
                continue
            body = code_only_html(rec["first"], claim["code"], claimed)
        else:
            body = body_html(rec["first"], claim["code"] if claim else None, claimed, p, args.as_of)
        rows.append({
            "Email": email, "FirstName": rec["first"], "LastName": rec["last"],
            "Campus": rec["campus"], "Period": rec["period"],
            "Claimed": "Y" if claimed else "N", "AccessCode": claim["code"] if claim else "",
            "DoneCount": p["done"], "TotalCount": p["total"], "PastDueCount": len(p["past_due"]),
            "FRQ2Status": p["frq_status"], "Subject": CODE_ONLY_SUBJECT if args.code_only else SUBJECT, "BodyHTML": body, "BodyText": to_text(body),
        })
    if args.code_only:
        for r in rows:
            for k in ("DoneCount", "TotalCount", "PastDueCount", "FRQ2Status"):
                r.pop(k)
    rows.sort(key=lambda r: (r["Campus"], ck.period_key({"period": r["Period"]}), r["LastName"], r["FirstName"]))

    stamp = args.as_of.isoformat() + ("_codeonly" if args.code_only else "")
    if args.test_to:
        rows = [dict(r, Email=args.test_to) for r in rows[:3]]
        stamp += "_TEST"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    base = OUT_DIR / f"unit2_email_merge_{stamp}"

    cols = list(rows[0].keys()) if rows else []
    with open(base.with_suffix(".csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)

    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb = Workbook()
    ws = wb.active
    ws.title = "Merge"
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    if rows:
        tab = Table(displayName="Merge", ref=f"A1:{get_column_letter(len(cols))}{len(rows) + 1}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
    wb.save(base.with_suffix(".xlsx"))

    preview = OUT_DIR / f"unit2_email_preview_{stamp}.html"
    sample = [r for r in rows if r["Claimed"] == "N"][:2] + [r for r in rows if r["Claimed"] == "Y"][:1]
    preview.write_text(
        "<!doctype html><meta charset='utf-8'><title>Email preview</title>"
        "<body style='font:14px Segoe UI,Arial;max-width:680px;margin:24px auto'>"
        + "".join(f"<div style='border:1px solid #ccc;border-radius:8px;padding:12px 18px;margin:0 0 18px'>"
                  f"<div style='color:#666;font-size:12px'>To: {esc(r['Email'])}<br>Subject: {esc(r['Subject'])}</div>"
                  f"<hr>{r['BodyHTML']}</div>" for r in sample) + "</body>", encoding="utf-8")

    n_un = sum(1 for r in rows if r["Claimed"] == "N")
    print(f"{len(rows)} emails ({n_un} with an access code to claim, {len(rows) - n_un} already set up)")
    print(f"  {base.with_suffix('.xlsx')}\n  {base.with_suffix('.csv')}\n  {preview}")
    if no_data:
        print(f"  NOTE: {len(no_data)} emailed students have no Unit 2 record (skipped): {no_data}")
    if no_code:
        print(f"  NOTE: {len(no_code)} students have no access code on file: {no_code}")


if __name__ == "__main__":
    main()
